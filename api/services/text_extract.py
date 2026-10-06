"""
Bytes → text for any document format, using only local tools (no API spend).

  pdf          → services.ocr_pipeline (text layer + Tesseract per page)
  docx         → python-docx (paragraphs + tables, in order)
  xlsx/xlsm    → openpyxl (each sheet as a markdown table)
  doc / xls    → antiword / LibreOffice headless when installed, else skipped
  images       → Tesseract
  txt/csv/html → decoded as text
Audio/video are handled by the meeting pipeline, not here.
"""
from __future__ import annotations

import io
import logging
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from typing import Optional

logger = logging.getLogger(__name__)

MEDIA_EXT = (".mp3", ".wav", ".wma", ".m4a", ".ogg", ".aac", ".flac",
             ".mp4", ".mov", ".m4v", ".avi", ".webm")


@dataclass
class ExtractResult:
    text: str = ""
    page_count: Optional[int] = None
    method: str = ""
    meta: dict = field(default_factory=dict)


def detect_kind(filename: str, content: bytes) -> str:
    fn = (filename or "").lower()
    for ext, kind in ((".pdf", "pdf"), (".docx", "docx"), (".doc", "doc"), (".xlsx", "xlsx"),
                      (".xlsm", "xlsx"), (".xls", "xls"), (".png", "image"), (".jpg", "image"),
                      (".jpeg", "image"), (".tif", "image"), (".tiff", "image"),
                      (".txt", "text"), (".csv", "text"), (".htm", "html"), (".html", "html")):
        if fn.endswith(ext):
            return kind
    if fn.endswith(MEDIA_EXT):
        return "media"
    head = content[:8]
    if head[:4] == b"%PDF":
        return "pdf"
    if head[:4] == b"PK\x03\x04":
        return "docx"
    if head == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "doc"
    if head[:3] == b"\xff\xd8\xff" or head == b"\x89PNG\r\n\x1a\n":
        return "image"
    return "unknown"


def _docx(content: bytes) -> str:
    from docx import Document as DocxDocument
    doc = DocxDocument(io.BytesIO(content))
    parts: list[str] = []
    for el in doc.element.body.iter():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            txt = "".join(t.text or "" for t in el.iter() if t.tag.split("}")[-1] == "t")
            if txt.strip():
                parts.append(txt)
        elif tag == "tbl":
            rows = []
            for tr in el.iter():
                if tr.tag.split("}")[-1] != "tr":
                    continue
                cells = ["".join(t.text or "" for t in tc.iter() if t.tag.split("}")[-1] == "t")
                         .replace("|", "\\|").strip()
                         for tc in tr.iter() if tc.tag.split("}")[-1] == "tc"]
                if any(cells):
                    rows.append("| " + " | ".join(cells) + " |")
            if rows:
                parts.append("\n".join(rows))
    return "\n\n".join(parts)


def _xlsx(content: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            if not any(c is not None and str(c).strip() for c in row):
                continue
            rows.append("| " + " | ".join("" if c is None else str(c).replace("|", "\\|").strip()
                                          for c in row) + " |")
        if rows:
            parts.append(f"## Sheet: {ws.title}\n\n" + "\n".join(rows))
    return "\n\n".join(parts)


def _run_tool(args: list[str], content: bytes, suffix: str, out_ext: Optional[str] = None) -> str:
    """Write content to a temp file, run a converter, return its text output."""
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, "in" + suffix)
        with open(src, "wb") as fh:
            fh.write(content)
        try:
            r = subprocess.run([*args, src] if not out_ext else [*args, "--outdir", td, src],
                               capture_output=True, timeout=120)
        except Exception as exc:
            logger.warning("%s failed: %s", args[0], exc)
            return ""
        if out_ext:
            out = os.path.join(td, "in" + out_ext)
            if os.path.exists(out):
                with open(out, "rb") as fh:
                    return fh.read().decode("utf-8", errors="replace")
            return ""
        return r.stdout.decode("utf-8", errors="replace") if r.returncode == 0 else ""


def _legacy_office(content: bytes, kind: str) -> tuple[str, str]:
    if kind == "doc" and shutil.which("antiword"):
        text = _run_tool(["antiword"], content, ".doc")
        if len(text.strip()) > 50:
            return text, "antiword"
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        target = "txt:Text" if kind == "doc" else "csv"
        text = _run_tool([soffice, "--headless", "--convert-to", target], content,
                         "." + kind, ".txt" if kind == "doc" else ".csv")
        if text.strip():
            return text, "libreoffice"
    return "", f"{kind}_unsupported"


def _image(content: bytes) -> str:
    from services.tesseract_ocr import is_tesseract_available
    if not is_tesseract_available():
        return ""
    import pytesseract
    from PIL import Image
    return pytesseract.image_to_string(Image.open(io.BytesIO(content)), config="--oem 1 --psm 3")


async def extract(filename: str, content: bytes) -> ExtractResult:
    kind = detect_kind(filename, content)
    if kind == "pdf":
        from services.ocr_pipeline import extract_pdf_to_markdown
        r = await extract_pdf_to_markdown(content, filename=filename)
        return ExtractResult(r.markdown, r.page_count, f"pdf:{r.tier}", {
            "ocr_pages": r.ocr_pages,
            "unreadable_pages": r.unreadable_pages[:50],
            "ocr_truncated": r.truncated,
        })
    if kind == "media":
        return ExtractResult(method="media_skipped")
    try:
        if kind == "docx":
            return ExtractResult(_docx(content), method="docx")
        if kind == "xlsx":
            return ExtractResult(_xlsx(content), method="xlsx")
        if kind in ("doc", "xls"):
            text, method = _legacy_office(content, kind)
            return ExtractResult(text, method=method)
        if kind == "image":
            return ExtractResult(_image(content), page_count=1, method="image:tesseract")
        if kind in ("text", "html"):
            text = content.decode("utf-8", errors="replace")
            if kind == "html":
                text = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", text, flags=re.S | re.I)
                text = re.sub(r"<[^>]+>", " ", text)
                text = re.sub(r"[ \t]+", " ", text)
            return ExtractResult(text, method=kind)
    except Exception as exc:
        logger.warning("extract %s (%s) failed: %s", filename, kind, exc)
        return ExtractResult(method=f"{kind}_error", meta={"extract_error": str(exc)[:300]})
    return ExtractResult(method="unknown_skipped")
